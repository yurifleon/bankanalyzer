import os
import tempfile
import unittest
from datetime import date
from decimal import Decimal
import bank_csv_monthly_dual_profile_cardnum as analyzer


def _make_tx(vendor, tx_date, debit=Decimal("0"), credit=Decimal("0"), description=None, card_number="", category=None):
    tx = {
        "row_number": 1,
        "date": tx_date,
        "month": analyzer.month_key(tx_date),
        "description": description or vendor,
        "vendor": vendor,
        "card_number": card_number,
        "debit": debit,
        "credit": credit,
        "net": credit - debit,
        "row": [],
    }

    # Left out entirely unless asked for, so tests can exercise the
    # Uncategorized fallback for transaction dicts built without a category.
    if category is not None:
        tx["category"] = category

    return tx


class TestBankCSVMonthlyDualProfile(unittest.TestCase):
    def test_parse_amount(self):
        self.assertEqual(analyzer.parse_amount("$1,234.56"), Decimal("1234.56"))
        self.assertEqual(analyzer.parse_amount("(1,234.56)"), Decimal("-1234.56"))
        self.assertEqual(analyzer.parse_amount(""), Decimal("0"))
        self.assertEqual(analyzer.parse_amount(None), Decimal("0"))

    def test_parse_date(self):
        self.assertEqual(analyzer.parse_date("06/08/2026").isoformat(), "2026-06-08")
        self.assertEqual(analyzer.parse_date("2026-06-08").isoformat(), "2026-06-08")
        self.assertEqual(analyzer.parse_date("06-08-26").isoformat(), "2026-06-08")
        self.assertIsNone(analyzer.parse_date("2026.06.08"))

    def test_clean_vendor_name(self):
        self.assertEqual(
            analyzer.clean_vendor_name("GglPay PANERA BREAD PENSACOLA  FL"),
            "PANERA BREAD"
        )
        self.assertEqual(
            analyzer.clean_vendor_name("AMAZON.COM/BILL"),
            "AMAZON"
        )
        self.assertEqual(
            analyzer.clean_vendor_name("SHELL SERVICE STATIOSUNRISE FL"),
            "SHELL SERVICE STATIOSUNRISE"
        )

    def test_safe_filename(self):
        self.assertEqual(analyzer.safe_filename("AMAZON MARKETPLACE"), "AMAZON_MARKETPLACE")
        self.assertEqual(analyzer.safe_filename("///"), "search")

    def test_detect_recurring_activity_finds_fixed_monthly_charge(self):
        transactions = [
            _make_tx("NETFLIX", date(2025, 1, 5), debit=Decimal("15.99")),
            _make_tx("NETFLIX", date(2025, 2, 5), debit=Decimal("15.99")),
            _make_tx("NETFLIX", date(2025, 3, 6), debit=Decimal("15.99")),
            _make_tx("NETFLIX", date(2025, 4, 5), debit=Decimal("15.99")),
        ]

        results = analyzer.detect_recurring_activity(transactions)

        self.assertEqual(len(results), 1)
        item = results[0]
        self.assertEqual(item["vendor"], "NETFLIX")
        self.assertTrue(item["is_recurring"])
        self.assertEqual(item["classification"], "Monthly Fixed Amount")
        self.assertEqual(item["direction"], "Expense / Charge")
        self.assertEqual(item["count"], 4)

    def test_detect_recurring_activity_ignores_infrequent_and_irregular(self):
        transactions = [
            # Only two occurrences: below the default min_occurrences threshold.
            _make_tx("RARE VENDOR", date(2025, 1, 1), debit=Decimal("50.00")),
            _make_tx("RARE VENDOR", date(2025, 6, 1), debit=Decimal("50.00")),
            # Irregular gaps and amounts: should not be classified as recurring.
            _make_tx("RANDOM SHOP", date(2025, 1, 3), debit=Decimal("12.00")),
            _make_tx("RANDOM SHOP", date(2025, 1, 20), debit=Decimal("87.00")),
            _make_tx("RANDOM SHOP", date(2025, 3, 15), debit=Decimal("5.00")),
        ]

        results = analyzer.detect_recurring_activity(transactions)

        by_vendor = {item["vendor"]: item for item in results}
        self.assertNotIn("RARE VENDOR", by_vendor)
        self.assertIn("RANDOM SHOP", by_vendor)
        self.assertFalse(by_vendor["RANDOM SHOP"]["is_recurring"])
        self.assertEqual(by_vendor["RANDOM SHOP"]["classification"], "Irregular")


class TestCategorizeTransaction(unittest.TestCase):
    def test_matches_merchant_on_cleaned_vendor(self):
        self.assertEqual(
            analyzer.categorize_transaction("PUBLIX", "PUBLIX 1234 PENSACOLA FL", Decimal("-52.10")),
            "Groceries",
        )

    def test_matches_structural_rule_on_raw_description(self):
        # clean_vendor_name strips TRANSFER/ACH as noise words, so the rule
        # can only fire if the raw description is also searched.
        self.assertEqual(
            analyzer.categorize_transaction("SAVINGS", "ACH TRANSFER TO SAVINGS", Decimal("-500")),
            "Transfers",
        )

    def test_structural_rules_take_precedence_over_merchant_rules(self):
        self.assertEqual(
            analyzer.categorize_transaction("AMAZON", "PAYROLL DEP AMAZON COM INC", Decimal("2500")),
            "Income",
        )

    def test_unmatched_credit_falls_back_to_income(self):
        self.assertEqual(
            analyzer.categorize_transaction("ZZQQ WIDGET", "ZZQQ WIDGET CO", Decimal("140.00")),
            "Income",
        )

    def test_unmatched_debit_falls_back_to_uncategorized(self):
        self.assertEqual(
            analyzer.categorize_transaction("ZZQQ WIDGET", "ZZQQ WIDGET CO", Decimal("-140.00")),
            "Uncategorized",
        )

    def test_subscriptions_outrank_utilities_for_apple_bilinternet(self):
        # _VENDOR_REPLACEMENTS rewrites BILINTERNET -> INTERNET, which the
        # Utilities rule would otherwise claim. Subscriptions must win.
        vendor = analyzer.clean_vendor_name("APPLE.COM/BILINTERNET")
        self.assertEqual(
            analyzer.categorize_transaction(vendor, "APPLE.COM/BILINTERNET", Decimal("-9.99")),
            "Subscriptions",
        )

    def test_keywords_match_whole_words_only(self):
        # "T-MOBILE" contains the Fuel & Transport keyword "MOBIL", and Fuel is
        # evaluated before Utilities. Only whole-word matching keeps the phone
        # bill out of the gas station bucket.
        self.assertEqual(
            analyzer.categorize_transaction("T MOBILE", "T-MOBILE PCS SVC", Decimal("-85.00")),
            "Utilities",
        )


    def test_matches_real_world_descriptors(self):
        # Forms that actually appear on statements: the apostrophe in
        # MCDONALD'S splits the token, and Amazon bills as AMZN MKTP.
        for description, expected in [
            ("MCDONALD'S F2231 PACE FL", "Dining"),
            ("AMZN MKTP US*2H4KL9DR3", "Shopping"),
        ]:
            with self.subTest(description=description):
                vendor = analyzer.clean_vendor_name(description)
                self.assertEqual(
                    analyzer.categorize_transaction(vendor, description, Decimal("-20.00")),
                    expected,
                )
    def test_rental_cars_are_travel_not_housing(self):
        # "RENT A CAR" contains the Housing & Rent keyword "RENT".
        for description in [
            "ENTERPRISE RENT A CAR PENSACOLA FL",
            "HERTZ RENT A CAR 1234",
            "AVIS RENT A CAR",
            "BUDGET RENT A CAR",
        ]:
            with self.subTest(description=description):
                vendor = analyzer.clean_vendor_name(description)
                self.assertEqual(
                    analyzer.categorize_transaction(vendor, description, Decimal("-210.00")),
                    "Travel",
                )

    def test_actual_rent_is_still_housing(self):
        description = "RENT PAYMENT OAKWOOD APTS"
        self.assertEqual(
            analyzer.categorize_transaction(
                analyzer.clean_vendor_name(description), description, Decimal("-1400.00")
            ),
            "Housing & Rent",
        )

    def test_payment_processor_prefix_does_not_claim_the_merchant(self):
        # NOISE_WORDS strips processor prefixes from the vendor so charges group
        # with the real merchant; the category rules must not undo that by
        # matching the prefix in the raw description.
        for description, expected in [
            ("GOOGLE *AMAZON MKTP", "Shopping"),
            ("GOOGLE *TARGET STORE", "Shopping"),
        ]:
            with self.subTest(description=description):
                vendor = analyzer.clean_vendor_name(description)
                self.assertEqual(
                    analyzer.categorize_transaction(vendor, description, Decimal("-30.00")),
                    expected,
                )

    def test_card_payments_are_transfers_not_income(self):
        # On the credit-card profile these are the largest credits on the
        # statement; the net>0 fallback must not report them as income.
        for description in [
            "PAYMENT RECEIVED - THANK YOU",
            "AUTOPAY 1234 - THANK YOU",
            "CHASE CREDIT CRD AUTOPAY",
            "MOBILE PAYMENT - THANK YOU",
        ]:
            with self.subTest(description=description):
                vendor = analyzer.clean_vendor_name(description)
                self.assertEqual(
                    analyzer.categorize_transaction(vendor, description, Decimal("1200.00")),
                    "Transfers",
                )

    def test_autopay_merchant_charges_are_not_swallowed_by_transfers(self):
        description = "STATE FARM INSURANCE AUTOPAY"
        self.assertEqual(
            analyzer.categorize_transaction(
                analyzer.clean_vendor_name(description), description, Decimal("-142.00")
            ),
            "Insurance",
        )

    def test_refund_credits_are_not_reported_as_income(self):
        # A refund is money coming back from a purchase. When the merchant is
        # recognizable it keeps that merchant's category; otherwise it falls to
        # Uncategorized. Either way it must not be counted as income.
        for description, expected in [
            ("RETURNED MERCHANDISE CREDIT", analyzer.UNCATEGORIZED),
            ("REFUND FROM WAYFAIR", "Shopping"),
        ]:
            with self.subTest(description=description):
                vendor = analyzer.clean_vendor_name(description)
                self.assertEqual(
                    analyzer.categorize_transaction(vendor, description, Decimal("64.00")),
                    expected,
                )


class TestSummarizeByMonthCategory(unittest.TestCase):
    def test_aggregates_totals_per_month_and_category(self):
        transactions = [
            _make_tx("PUBLIX", date(2026, 3, 4), debit=Decimal("50.00"), category="Groceries"),
            _make_tx("KROGER", date(2026, 3, 9), debit=Decimal("25.50"), category="Groceries"),
            _make_tx("EMPLOYER", date(2026, 3, 1), credit=Decimal("2000.00"), category="Income"),
        ]

        rows = analyzer.summarize_by_month_category(transactions)
        groceries = [r for r in rows if r["category"] == "Groceries"][0]

        self.assertEqual(groceries["month"], "2026-03")
        self.assertEqual(groceries["count"], 2)
        self.assertEqual(groceries["total_debit"], Decimal("75.50"))
        self.assertEqual(groceries["total_credit"], Decimal("0"))

    def test_transactions_without_a_category_key_fall_back_to_uncategorized(self):
        rows = analyzer.summarize_by_month_category(
            [_make_tx("ZZQQ", date(2026, 3, 4), debit=Decimal("10.00"))]
        )

        self.assertEqual(rows[0]["category"], analyzer.UNCATEGORIZED)

    def test_sorted_by_month_then_category(self):
        transactions = [
            _make_tx("X", date(2026, 4, 1), debit=Decimal("1"), category="Utilities"),
            _make_tx("Y", date(2026, 3, 1), debit=Decimal("1"), category="Shopping"),
            _make_tx("Z", date(2026, 3, 1), debit=Decimal("1"), category="Dining"),
        ]

        rows = analyzer.summarize_by_month_category(transactions)

        self.assertEqual(
            [(r["month"], r["category"]) for r in rows],
            [("2026-03", "Dining"), ("2026-03", "Shopping"), ("2026-04", "Utilities")],
        )


class TestReadTransactionsCategory(unittest.TestCase):
    def test_read_transactions_assigns_a_category_to_each_row(self):
        handle, path = tempfile.mkstemp(suffix=".csv")
        os.close(handle)

        try:
            with open(path, "w", encoding="utf-8", newline="") as fh:
                fh.write("ref,date,x,description,debit,credit\n")
                fh.write("1,03/04/2026,,PUBLIX 1234 PENSACOLA FL,52.10,\n")
                fh.write("2,03/05/2026,,ZZQQ WIDGET CO,140.00,\n")

            transactions = analyzer.read_transactions(
                path, date_col=1, text_col=3, debit_col=4, credit_col=5
            )

            self.assertEqual(
                [tx["category"] for tx in transactions],
                ["Groceries", analyzer.UNCATEGORIZED],
            )
        finally:
            os.unlink(path)


class TestCategorySheet(unittest.TestCase):
    def test_workbook_contains_a_monthly_by_category_sheet(self):
        from openpyxl import load_workbook

        transactions = [
            _make_tx("PUBLIX", date(2026, 3, 4), debit=Decimal("50.00"), category="Groceries"),
            _make_tx("KROGER", date(2026, 3, 9), debit=Decimal("25.50"), category="Groceries"),
            _make_tx("EMPLOYER", date(2026, 3, 1), credit=Decimal("2000.00"), category="Income"),
        ]

        handle, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(handle)

        try:
            analyzer.write_workbook(transactions, path)
            wb = load_workbook(path)

            self.assertIn("Monthly by Category", wb.sheetnames)

            rows = list(wb["Monthly by Category"].values)
            self.assertEqual(
                rows[0],
                ("Month", "Category", "Transaction Count",
                 "Total Debit / Charges", "Total Credit / Payments", "Net Credit - Debit"),
            )
            self.assertEqual(rows[1], ("2026-03", "Groceries", 2, 75.50, 0.0, -75.50))
            self.assertEqual(rows[2], ("2026-03", "Income", 1, 0.0, 2000.0, 2000.0))
        finally:
            os.unlink(path)
